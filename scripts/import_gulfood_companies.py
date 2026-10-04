"""
ورود شرکت‌های فایل اکسل نمایشگاه (خروجی exhibitor_lead_finder.py، مثل «Gulfood 2026.xlsx»)
به بانک شرکت‌ها (data/companies.json).

قاعده‌ی کاربر (۲۰۲۶-۱۰-۰۴): فقط شرکت‌هایی که هم وب‌سایت و هم ایمیل دارند وارد می‌شن؛
بدون ایمیل (یا بدون وب‌سایت) وارد نمی‌شه. علاوه‌بر اون، چون «وب‌سایت رسمی» لازمه:
  - وب‌سایتِ حدسی (منبع «جست‌وجوی وب (حدسی)») و لینک‌های شبکه‌ی اجتماعی/تجمیع‌کننده
    (linkedin، linktr.ee، خود gulfood.com، ...) به‌جای وب‌سایت رسمی پذیرفته نمی‌شن.
  - ایمیل از extract_email (scripts/clean_company_emails.py) رد می‌شه — همون قاعده‌ی
    «بعد از هر ایمپورت لیست شرکت، ایمیل رو پاک کن».
  - شرکتی که ایمیل یا دامنه‌ی وب‌سایتش یا نامش از قبل توی companies.json هست، دوباره
    اضافه نمی‌شه (و داخل خود فایل هم تکراریِ هم‌نام+هم‌ایمیل حذف می‌شه؛ شعبه‌ها با ایمیل
    مشترک ولی نام متفاوت مشکلی ندارن).

اجرا:
    python3 scripts/import_gulfood_companies.py "<مسیر xlsx>" --event "Gulfood 2026"          # dry-run
    python3 scripts/import_gulfood_companies.py "<مسیر xlsx>" --event "Gulfood 2026" --apply
"""

import argparse
import json
import os
import re
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from clean_company_emails import extract_email  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
COMPANIES_FILE = os.path.join(BASE_DIR, "data", "companies.json")
COUNTRY_MAP_FILE = os.path.join(BASE_DIR, "scripts", "country_name_map.json")

# نام‌هایی که توی نقشه‌ی استاندارد (country_name_map.json) با همین املا نیستن.
COUNTRY_ALIASES = {
    "iran": "ایران",
    "vietnam": "ویتنام",
    "republic of korea": "کره جنوبی",
    "taiwan province, people’s republic of china": "تایوان",
    "republic of ireland": "ایرلند",
    "syria": "سوریه",
    "palestine": "فلسطین",
    "kosovo": "کوزوو",
    "libya": "لیبی",
    "hong kong special administrative region": "هنگ‌کنگ",
    "monaco": "موناکو",
    "tanzania": "تانزانیا",
    "wallis & futuna": "والیس و فوتونا",
}

# وب‌سایت‌هایی که «وب‌سایت رسمی شرکت» نیستن.
NON_OFFICIAL_HOSTS = (
    "companywebsite.com", "example.com", "facebook.com", "instagram.com", "linkedin.com",
    "twitter.com", "x.com", "youtube.com", "wa.me", "t.me", "tiktok.com", "gulfood.com",
    "linktr.ee", "pinterest.com",
)
GUESSED_SOURCE_MARKER = "حدسی"


def host_of(url):
    return re.sub(r"^https?://(www\.)?", "", str(url).strip().lower()).split("/")[0].split("?")[0]


def official_site(url, source):
    if not url:
        return None
    url = str(url).strip()
    if source and GUESSED_SOURCE_MARKER in str(source):
        return None
    candidate = url if url.lower().startswith("http") else "https://" + url
    host = host_of(candidate)
    if not re.match(r"^[a-z0-9.\-]+\.[a-z]{2,}$", host):
        return None
    if any(host == b or host.endswith("." + b) for b in NON_OFFICIAL_HOSTS):
        return None
    return url


def norm_name(s):
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def normalize_grade(text):
    # classifyGrade (web/lib/outreachTemplate.js) فقط «خوراک» رو به‌عنوان گرید خوراکی
    # می‌شناسه، نه «غذایی» — پس اینجا یکدست می‌کنیم تا ایمیل معرفی گرید درست رو تشخیص بده.
    t = (text or "").replace("گرید غذایی", "خوراکی").replace("غذایی", "خوراکی").strip()
    return t or "خوراکی"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--event", required=True, help="نام نمایشگاه، مثلاً «Gulfood 2026»")
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()

    country_map = {k.lower(): v["fa"] for k, v in json.load(open(COUNTRY_MAP_FILE, encoding="utf-8")).items()}
    country_map.update(COUNTRY_ALIASES)

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    ws = wb.worksheets[0]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[1]]

    companies = json.load(open(COMPANIES_FILE, encoding="utf-8"))
    ex_emails = {extract_email(c.get("email"))[0] for c in companies} - {None}
    ex_hosts = {host_of(c["website"]) for c in companies if c.get("website")}
    ex_names = {norm_name(c.get("english_name")) for c in companies}
    next_id = max(int(c["id"].split("-")[1]) for c in companies) + 1

    stats = {"rows": len(rows), "no_email": 0, "no_official_site": 0, "dup_existing": 0, "dup_in_file": 0,
             "unmapped_country": {}}
    new, seen = [], set()
    for r in rows:
        email, _ = extract_email(r.get("ایمیل"))
        if not email:
            stats["no_email"] += 1
            continue
        site = official_site(r.get("وب‌سایت"), r.get("منبع وب‌سایت"))
        if not site:
            stats["no_official_site"] += 1
            continue
        name = str(r["نام شرکت"]).strip()
        if email in ex_emails or host_of(site) in ex_hosts or norm_name(name) in ex_names:
            stats["dup_existing"] += 1
            continue
        key = (norm_name(name), email)
        if key in seen:
            stats["dup_in_file"] += 1
            continue
        seen.add(key)

        en_country = (r.get("کشور") or "").strip()
        fa_country = country_map.get(en_country.lower())
        if not fa_country:
            stats["unmapped_country"][en_country] = stats["unmapped_country"].get(en_country, 0) + 1
            fa_country = en_country  # حذف نمی‌کنیم؛ با نام انگلیسی ثبت می‌شه
        kind = r.get("نوع")
        categories = [c.strip() for c in str(r.get("دسته‌بندی‌های سایت") or "").split(",") if c.strip()][:4]
        industry = " — ".join(x for x in (kind, "، ".join(categories)) if x) or "شرکت غذایی/نوشیدنی"
        stand = r.get("غرفه")
        profile = r.get("لینک پروفایل")
        note_parts = [f"نمایشگاه {args.event}"]
        if stand:
            note_parts.append(str(stand))
        if profile:
            note_parts.append(str(profile))
        record = {
            "country": fa_country,
            "english_name": name,
            "industry": industry,
            "target_grade": normalize_grade(r.get("گرید پیشنهادی")),
            "phone": str(r["واتساپ/تلفن"]).strip() if r.get("واتساپ/تلفن") else None,
            "email": email,
            "website": site,
            "source_note": " · ".join(note_parts),
            "id": f"company-{next_id}",
        }
        record = {k: v for k, v in record.items() if v is not None}
        next_id += 1
        new.append(record)

    print(f"ردیف‌های فایل: {stats['rows']}")
    print(f"  بدون ایمیل معتبر (وارد نمی‌شن): {stats['no_email']}")
    print(f"  بدون وب‌سایت رسمی: {stats['no_official_site']}")
    print(f"  تکراری با بانک فعلی: {stats['dup_existing']}  | تکراری داخل فایل: {stats['dup_in_file']}")
    print(f"  کشور بدون نگاشت فارسی: {stats['unmapped_country']}")
    print(f"=> شرکت جدید برای ورود: {len(new)}  (id از {new[0]['id'] if new else '-'} تا {new[-1]['id'] if new else '-'})")

    if not args.apply:
        print("\n[dry-run] چیزی نوشته نشد. با --apply واقعاً وارد می‌شه.")
        return
    companies.extend(new)
    with open(COMPANIES_FILE, "w", encoding="utf-8") as f:
        json.dump(companies, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print(f"\n[DONE] {len(new)} شرکت به {COMPANIES_FILE} اضافه شد (کل: {len(companies)}).")


if __name__ == "__main__":
    main()
