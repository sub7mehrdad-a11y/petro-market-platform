"""
اکسل نهایی «فقط اطلاعات تأییدشده» برای یک نمایشگاه.

ورودی: data/exhibition_leads/<slug>.progress.json (شرکت‌های مرتبط) و <slug>.verified.json
(خروجی scripts/verify_company_contacts.py). هر مقدار تماس که در verified نیست، در اکسل نمی‌آید —
عمداً هیچ ایمیل/تلفن/وب‌سایتِ تأییدنشده‌ای از progress کپی نمی‌شود.

استفاده:  python scripts/export_verified_xlsx.py <slug> "<نام نمایشگاه>"
"""

import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(__file__))
import exhibitor_lead_finder as efl  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

EMAIL_CONF_FA = {
    "site_confirmed_same_domain": "تأییدشده روی سایت شرکت",
    "site_confirmed_other_domain": "روی سایت شرکت دیده شد (دامنه‌ی ایمیل متفاوت)",
    "directory_only": "ثبت‌شده توسط نمایشگاه — روی سایت دیده نشد (دامنه معتبر)",
}
SITE_FA = {"verified_official": "فهرست رسمی نمایشگاه + باز می‌شود",
           "verified_name_match": "باز می‌شود + نام شرکت روی صفحه تأیید شد"}
COMPANY_TYPE_FA = efl.COMPANY_TYPE_FA


def main():
    slug, event = sys.argv[1], sys.argv[2]
    prog = efl.load_progress(os.path.join(efl.OUTPUT_DIR, f"{slug}.progress.json"))
    ver = efl.load_progress(os.path.join(efl.OUTPUT_DIR, f"{slug}.verified.json"))
    companies = [(k, v) for k, v in prog.items() if v.get("relevant")]
    companies.sort(key=lambda kv: kv[1].get("_order", 0))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "شرکت‌ها"
    ws.sheet_view.rightToLeft = True
    headers = ["ردیف", "نام شرکت", "کشور", "نوع", "گرید پیشنهادی", "وب‌سایت (تأییدشده)", "وضعیت وب‌سایت",
               "ایمیل (تأییدشده)", "اعتبار ایمیل", "صفحه‌ی منبع ایمیل", "تلفن/واتساپ (معتبر)", "صفحه‌ی منبع تلفن",
               "وضعیت تماس", "نیاز به بررسی دستی؟", "دلیل (مدل)", "توضیحات شرکت", "غرفه", "شهر/آدرس"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")

    stats = {"total": 0, "email": 0, "email_site": 0, "phone": 0, "site": 0, "unverified": 0}
    for i, (k, c) in enumerate(companies, start=1):
        v = ver.get(k, {})
        email, econf = v.get("email_best"), v.get("email_confidence")
        phone = v.get("phone_best")
        site = v.get("website")
        status_site = SITE_FA.get(v.get("website_status"), "")
        weak = []
        if v.get("website_match") == "medium":
            weak.append("تطبیق نام سایت ضعیف‌تر")
        if econf == "site_confirmed_other_domain":
            weak.append("دامنه‌ی ایمیل با سایت فرق دارد")
        if c.get("needs_review"):
            weak.append("طبقه‌بندی شرکت نیاز به بررسی دارد")
        if email and phone:
            st = "ایمیل + تلفن"
        elif email:
            st = "فقط ایمیل"
        elif phone:
            st = "فقط تلفن"
        else:
            st = "هیچ تماس تأییدشده‌ای نیست"
        ws.append([
            i, c.get("name", ""), c.get("country", ""), COMPANY_TYPE_FA.get(c.get("company_type"), c.get("company_type") or ""),
            c.get("suggested_grade", ""), site or "", status_site if site else "",
            email or "", EMAIL_CONF_FA.get(econf, "") if email else "",
            (v.get("email_source_url") or "") if email else "", phone or "", (v.get("phone_source_url") or "") if phone else "",
            st, "؛ ".join(weak), c.get("llm_reason", ""), (c.get("description") or "")[:300], c.get("stand", ""),
            c.get("region_address", ""),
        ])
        stats["total"] += 1
        stats["email"] += bool(email)
        stats["email_site"] += bool(email and str(econf).startswith("site"))
        stats["phone"] += bool(phone)
        stats["site"] += bool(site)
        stats["unverified"] += (k not in ver)

    widths = [6, 32, 14, 16, 30, 30, 26, 30, 34, 44, 20, 40, 22, 28, 40, 55, 20, 30]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"

    out = os.path.join(efl.EXTERNAL_XLSX_DIR, f"{efl.safe_filename(event)} (تأییدشده).xlsx")
    tmp = out + ".tmp"
    wb.save(tmp)
    os.replace(tmp, out)
    print(f"{out}\n{stats}")


if __name__ == "__main__":
    main()
