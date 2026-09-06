"""
ورود فایل «ایمیل‌های ارسال‌شده.xlsx» (کمپین قبلی ایمیل معرفی، عراق - ۲۶ جولای) به
data/email_outreach_sent.json — تا سیستم ایمیل معرفی جدید هیچ‌وقت به همین
آدرس‌ها دوباره ایمیل نفرسته.

چرا این فایل جدا از data/companies.json است: آدرس‌های ارسال‌شده لزوماً با
شرکت‌های فعلی دیتابیس یک‌به‌یک نیستن (۷۵ آدرس قبلی، فقط ۲۸ تاش توی
companies.json فعلی پیدا شد) — پس باید یک فهرست مستقل نگه داریم و همیشه
بر اساس خودِ آدرس ایمیل (نه id شرکت) چک کنیم.

اجرای دوباره این اسکریپت، رکوردهای منبع «legacy-iraq-campaign» رو جایگزین
می‌کنه (نه اضافه)؛ رکوردهایی که بعداً توسط خودِ سیستم ارسال جدید اضافه شدن
(source متفاوت) دست‌نخورده می‌مونن.
"""

import os
import json

import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(__file__))
SRC_FILE = os.path.join(BASE_DIR, "ایمیل بازاریابی", "ایمیل های ارسال شده.xlsx")
OUTPUT_FILE = os.path.join(BASE_DIR, "data", "email_outreach_sent.json")
SOURCE_TAG = "legacy-iraq-campaign"


def parse_sent_list(path: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    records = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if not row or not row[2]:
            continue
        name, phone, email, sent_at = row[0], row[1], row[2], row[3]
        records.append({
            "email": str(email).strip().lower(),
            "name": str(name).strip() if name else None,
            "phone": str(phone).strip() if phone is not None else None,
            "sent_at": sent_at.isoformat() if hasattr(sent_at, "isoformat") else str(sent_at),
            "source": SOURCE_TAG,
        })
    return records


def main():
    if not os.path.exists(SRC_FILE):
        print(f"[WARN] فایل پیدا نشد: {SRC_FILE}")
        return

    legacy_records = parse_sent_list(SRC_FILE)

    existing = []
    if os.path.exists(OUTPUT_FILE):
        with open(OUTPUT_FILE, "r", encoding="utf-8") as f:
            existing = json.load(f)
    kept = [r for r in existing if r.get("source") != SOURCE_TAG]

    combined = kept + legacy_records
    os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        json.dump(combined, f, ensure_ascii=False, indent=2)

    print(f"[OK] {len(legacy_records)} آدرس از کمپین قدیمی + {len(kept)} رکورد دیگر = "
          f"{len(combined)} رکورد کل در {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
