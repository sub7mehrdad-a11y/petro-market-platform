"""
ورود لیست شرکت‌های هدف و نمایشگاه‌های عربستان (پوشه‌ی «تحلیل بازار عربستان»، مهر ۱۴۰۵) به پلتفرم.

شرکت‌ها (لیست_شرکتهای_هدف_عربستان_1.xlsx — ۳۰ شرکت با تماس تأییدشده از سایت رسمی):
  - ۵ شرکت از قبل (از فهرست Gulfood 2026) در بانک بودند → فیلدهای تماس/گرید/پتانسیل با نسخه‌ی تأییدشده به‌روز می‌شوند؛
    اگر ایمیل عوض شد، ایمیل قبلی در email_history می‌ماند (OVERLAPS: نگاشت دستی چون دامنه‌ها فرق دارند).
  - بقیه شرکت‌های جدید با شناسه‌ی بعدی اضافه می‌شوند (ایمیل با extract_email تمیز می‌شود).
نمایشگاه‌ها (لیست_نمایشگاههای_تجاری_عربستان.xlsx): اگر هم‌نام و هم‌کشور از قبل بود (Gulfood Manufacturing) دست نمی‌خورد؛
  بقیه با شناسه‌ی بعدی اضافه می‌شوند.

اجرا:
    python scripts/import_saudi_lists.py            # dry-run
    python scripts/import_saudi_lists.py --apply
"""

import argparse
import datetime
import json
import os
import re
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from clean_company_emails import extract_email  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
COMPANIES_FILE = os.path.join(BASE_DIR, "data", "companies.json")
EXHIBITIONS_FILE = os.path.join(BASE_DIR, "data", "exhibitions.json")
SRC_DIR = r"E:\مهرداد\سپهران\تحلیل بازار عربستان"
COMPANIES_XLSX = os.path.join(SRC_DIR, "لیست_شرکتهای_هدف_عربستان_1.xlsx")
EXHIBITIONS_XLSX = os.path.join(SRC_DIR, "لیست_نمایشگاههای_تجاری_عربستان.xlsx")
COUNTRY = "عربستان سعودی"
SOURCE_FILE = "لیست_شرکتهای_هدف_عربستان_1.xlsx"

# نام (کوچک‌شده‌ی شروع) در فایل جدید → شناسه‌ی رکورد موجود در بانک
OVERLAPS = {
    "riyadh food industries": "company-1658",
    "ahmad a. abed": "company-598",
    "national biscuits and confectionery": "company-1424",
    "sunbulah group": "company-845",
    "savola group": "company-1684",
}


def norm_grade(text):
    t = (text or "").strip()
    return t.replace("خوراک دام", "دام و طیور") or "نامشخص"


def build_record(row, today):
    name, local, activity, grade, potential, city, address, phone, email, site, how = (row + (None,) * 11)[:11]
    addr = " — ".join(x for x in (str(city or "").strip(), str(address or "").strip()) if x and x != "نامشخص")
    rec = {
        "country": COUNTRY,
        "english_name": str(name).strip(),
        "industry": str(activity or "").strip() or None,
        "target_grade": norm_grade(grade),
        "purchasing_potential": str(potential or "").strip() or None,
        "address": addr or None,
        "phone": str(phone).strip() if phone else None,
        "email": extract_email(email)[0],
        "website": str(site).strip() if site else None,
        "source_note": f"لیست شرکت‌های هدف عربستان (گزارش مهر ۱۴۰۵) · تماس تأییدشده از سایت رسمی · {str(how or '').strip()}".strip(" ·"),
    }
    if local:
        rec["arabic_name"] = str(local).strip()
    return {k: v for k, v in rec.items() if v}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()
    today = datetime.date.today().isoformat()

    companies = json.load(open(COMPANIES_FILE, encoding="utf-8"))
    by_id = {c["id"]: c for c in companies}
    next_id = max(int(c["id"].split("-")[1]) for c in companies) + 1
    used_emails = {(extract_email(c.get("email"))[0] or "").lower(): c["id"] for c in companies}

    ws = openpyxl.load_workbook(COMPANIES_XLSX, data_only=True).worksheets[0]
    rows = [tuple(r) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    added, updated, skipped = [], [], []
    for row in rows:
        rec = build_record(row, today)
        if not rec.get("email"):
            skipped.append((rec["english_name"], "ایمیل نامعتبر"))
            continue
        key = rec["english_name"].lower()
        match_id = next((cid for prefix, cid in OVERLAPS.items() if key.startswith(prefix)), None)
        if match_id:
            old = by_id[match_id]
            changes = []
            for field in ("email", "website", "phone", "address", "purchasing_potential", "target_grade", "industry", "arabic_name"):
                if rec.get(field) and rec[field] != old.get(field):
                    if field == "email":
                        old.setdefault("email_history", []).append({"email": old.get("email"), "reason": "جایگزینی با ایمیل تأییدشده‌ی لیست عربستان", "date": today, "action": "replace", "replaced_with": rec["email"]})
                    changes.append(field)
                    old[field] = rec[field]
            if changes:
                old["source_note"] = (old.get("source_note", "") + " · به‌روزرسانی با لیست شرکت‌های هدف عربستان (مهر ۱۴۰۵)").strip(" ·")
            updated.append((match_id, rec["english_name"], changes))
            continue
        owner = used_emails.get(rec["email"].lower())
        if owner:
            skipped.append((rec["english_name"], f"ایمیل قبلاً برای {owner} ثبت است"))
            continue
        rec["id"] = f"company-{next_id}"
        next_id += 1
        used_emails[rec["email"].lower()] = rec["id"]
        added.append(rec)
        companies.append(rec)

    exhibitions = json.load(open(EXHIBITIONS_FILE, encoding="utf-8"))
    ex_keys = {(e["name"].strip().lower(), e["country"]) for e in exhibitions}
    ex_next = max(int(e["id"].split("-")[1]) for e in exhibitions) + 1
    ws = openpyxl.load_workbook(EXHIBITIONS_XLSX, data_only=True).worksheets[0]
    ex_added, ex_skipped = [], []
    for r in list(ws.iter_rows(min_row=2, values_only=True)):
        if not r[0]:
            continue
        name, focus, place, timing, site = (tuple(r) + (None,) * 5)[:5]
        place = str(place or "")
        country = "امارات" if "امارات" in place else COUNTRY
        if (str(name).strip().lower(), country) in ex_keys:
            ex_skipped.append(name)
            continue
        rec = {"country": country, "source_file": os.path.basename(EXHIBITIONS_XLSX), "name": str(name).strip(),
               "focus": str(focus or "").strip(), "location": place.replace(" / ", "، "), "date": str(timing or "").strip(),
               "id": f"exhibition-{ex_next}"}
        if site:
            rec["website"] = str(site).strip()
        ex_next += 1
        ex_added.append(rec)
        exhibitions.append(rec)

    print(f"شرکت‌ها: ردیف فایل {len(rows)} | جدید {len(added)} | به‌روزشده {len(updated)} | ردشده {len(skipped)}")
    for cid, n, ch in updated:
        print(f"   به‌روز: {cid} {n}: {', '.join(ch) or 'بدون تغییر'}")
    for n, why in skipped:
        print(f"   رد: {n}: {why}")
    print(f"نمایشگاه‌ها: جدید {len(ex_added)} ({[e['name'] for e in ex_added]}) | تکراری (دست‌نخورد) {ex_skipped}")
    if not args.apply:
        print("\n[dry-run] چیزی نوشته نشد.")
        return
    for path, data in ((COMPANIES_FILE, companies), (EXHIBITIONS_FILE, exhibitions)):
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
            f.write("\n")
    print("\n[DONE] نوشته شد.")


if __name__ == "__main__":
    main()
